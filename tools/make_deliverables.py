# -*- coding: utf-8 -*-
"""生成课程提交所需的交付物。

把 Markdown 报告转成 Word(.docx)，生成分工表 Excel(.xlsx)，
并把 src/ 下的核心代码合并成单个 核心代码.py。

用法：
    python tools/make_deliverables.py
"""

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "报告"
OUT_DIR = ROOT / "交付物"


def _set_cjk_font(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "宋体")


def _add_md_paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    for part in re.split(r"(\*\*.*?\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            p.add_run(part.replace("`", ""))
    return p


def _add_table(doc, table_lines):
    rows = []
    for tl in table_lines:
        cells = [c.strip() for c in tl.strip().strip("|").split("|")]
        rows.append(cells)
    data_rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    if not data_rows:
        return
    ncols = max(len(r) for r in data_rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    for r in data_rows:
        cells = table.add_row().cells
        for j in range(ncols):
            cells[j].text = r[j] if j < len(r) else ""


def md_to_docx(md_path: Path, docx_path: Path):
    doc = Document()
    _set_cjk_font(doc)
    lines = md_path.read_text(encoding="utf-8").split("\n")
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()
        # 代码块
        if stripped.startswith("```"):
            code = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            if code:
                p = doc.add_paragraph()
                p.style = doc.styles["No Spacing"]
                run = p.add_run("\n".join(code))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            continue
        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
            i += 1
            continue
        # 表格
        if stripped.startswith("|"):
            tbl = []
            while i < n and lines[i].strip().startswith("|"):
                tbl.append(lines[i])
                i += 1
            _add_table(doc, tbl)
            continue
        # 无序列表
        if stripped.startswith("- "):
            _add_md_paragraph(doc, stripped[2:], style="List Bullet")
            i += 1
            continue
        # 引用
        if stripped.startswith("> "):
            _add_md_paragraph(doc, stripped[2:])
            i += 1
            continue
        # 分隔线 / 空行
        if stripped in ("", "---"):
            i += 1
            continue
        # 普通段落
        _add_md_paragraph(doc, stripped)
        i += 1
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    print(f"  [docx] {docx_path.name}")


def make_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "分组情况及职责分配"
    ws.append(["组号", "成员姓名", "学号", "主要职责", "对应代码/文件"])
    rows = [
        ["第25组", "WJ（替换为真实姓名）", "（填写学号）", "数据预处理与环境搭建", "src/data.py、requirements.txt、README.md"],
        ["第25组", "WWD（替换为真实姓名）", "（填写学号）", "模型设计与训练", "src/model.py、src/train.py"],
        ["第25组", "CGS（替换为真实姓名）", "（填写学号）", "实验评估与报告撰写", "src/evaluate.py、报告/、PPT/"],
    ]
    for r in rows:
        ws.append(r)
    # 列宽
    for col, w in zip("ABCDE", [8, 24, 14, 26, 40]):
        ws.column_dimensions[col].width = w
    out = OUT_DIR / "分组情况及职责分配表.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"  [xlsx] {out.name}")


def make_core_code():
    parts = []
    header = "#" * 70
    for name in ["model.py", "data.py", "train.py", "evaluate.py"]:
        f = ROOT / "src" / name
        parts.append(f"{header}\n# 文件：src/{name}\n{header}\n")
        parts.append(f.read_text(encoding="utf-8").rstrip())
        parts.append("\n\n")
    out = OUT_DIR / "核心代码.py"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(parts), encoding="utf-8")
    print(f"  [py]   {out.name}")


def main():
    print("生成交付物 ->", OUT_DIR)
    md_files = {
        "汇总报告.md": "课程报告（汇总）.docx",
        "个人报告_WJ.md": "课程报告_WJ.docx",
        "个人报告_WWD.md": "课程报告_WWD.docx",
        "个人报告_CGS.md": "课程报告_CGS.docx",
    }
    for src, dst in md_files.items():
        md_to_docx(REPORT_DIR / src, OUT_DIR / dst)
    make_xlsx()
    make_core_code()
    print("全部生成完成。")


if __name__ == "__main__":
    main()
